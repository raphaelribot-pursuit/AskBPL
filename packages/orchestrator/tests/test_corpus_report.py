"""Patch 0020: `askbpl-corpus-report` and docs/CURATOR_WORKFLOW.md (PRD ING-13, ING-7, ING-9, AC-1, NFR-6).

The report is read-only evidence for people deciding what the chatbot may say. These tests exist to stop it from
(a) counting anything twice, (b) implying that more is ready than is, (c) writing anything, and (d) printing a secret.
The guide's tests keep every command and flag it names real, and pin the three behaviours it warns about.
"""

from __future__ import annotations

import asyncio
import csv
import re
import shlex
import sys
import uuid
from pathlib import Path
from types import SimpleNamespace
from typing import Any, Optional
from unittest.mock import patch

import openpyxl
import pytest
from sqlalchemy import delete, func, select

from ask_bpl_orchestrator import corpus_report as cr
from ask_bpl_orchestrator.config import Settings, get_settings
from ask_bpl_orchestrator.db.models import Chunk, Document, Eligibility, ExtractionStatus, IngestStatus, SourceFormat
from ask_bpl_orchestrator.db.session import async_session_factory
from ask_bpl_orchestrator.ingest.inventory import InventoryRow, apply_inventory_to_db, normalize_name

ORCH = Path(__file__).resolve().parents[1]
REPO = ORCH.parents[1]
INVENTORY = ORCH / "sample_corpus" / "AskBPL_Source_Inventory_v0_1.xlsx"

HEADER = [
    "Source ID", "File name", "Format", "Pages", "Chars per page", "Parse check", "In current pilot index",
    "Source type", "Lifecycle status (ING-9)", "Blocking decision", "Notes (PRD v1.2 s19.3 / 19.4)", "Owner (D-02)",
    "Policy number / dates (from document, to confirm)", "Reviewer / date",
]


# ------------------------------------------------------------------------------------------------ builders
def wrow(sid: str, name: str, *, parse: str = "Pass: text layer present", pilot: str = "No", typ: str = "Policy",
         life: str = "Authority pending", blocking: str = "", reviewer: Optional[str] = None) -> dict[str, Any]:
    return dict(sid=sid, name=name, parse=parse, pilot=pilot, typ=typ, life=life, blocking=blocking, reviewer=reviewer)


def make_workbook(path: Path, rows: list[dict[str, Any]], counts: Optional[dict[str, Any]] = None) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Source Inventory"
    ws.append(HEADER)
    for r in rows:
        ws.append([r["sid"], r["name"], "PDF", 3, 2000, r["parse"], r["pilot"], r["typ"], r["life"], r["blocking"],
                   None, None, None, r["reviewer"]])
    if counts is not None:
        cs = wb.create_sheet("Counts (ING-13)")
        cs.append(["Measure", "Count", "How it is computed"])
        for label, value in counts.items():
            cs.append([label, value, ""])
    wb.save(path)
    return path


def sources(tmp_path: Path, rows: list[dict[str, Any]], counts: Optional[dict[str, Any]] = None) -> list[cr.SourceRow]:
    return cr.read_sources(make_workbook(tmp_path / "inv.xlsx", rows, counts))


def doc(path: str, *, elig: str = "authority_pending", member: bool = True, status: str = "indexed", chunks: int = 3,
        answerable: bool = True, extraction: str = "standard", superseded: bool = False, sid: Optional[str] = None,
        did: Optional[str] = None) -> cr.DbDoc:
    return cr.DbDoc(did or str(uuid.uuid4()), "t", path, sid, elig, member, status, extraction, superseded, chunks, answerable)


def facts_with(*docs: cr.DbDoc, allow: bool = True) -> cr.CorpusFacts:
    return cr.CorpusFacts(state="checked", target="localhost:5433/x", allow_test_index=allow, docs=list(docs))


def counts_of(report: cr.CorpusReport) -> dict[str, str]:
    return {c.label.strip(): c.value for c in report.counts}


def codes(report: cr.CorpusReport, subject: Optional[str] = None) -> set[str]:
    return {f.code for f in report.findings if subject is None or f.subject == subject}


# ------------------------------------------------------------------------------------------------ counts
def test_shipped_inventory_buckets_add_up_and_nothing_is_double_counted() -> None:
    rows = cr.read_sources(INVENTORY)
    rep = cr.build_report(rows, cr.CorpusFacts(), counts_sheet=cr.read_counts_sheet(INVENTORY))
    c = counts_of(rep)
    assert c["Received (rows in the inventory)"] == "30"
    assert c["Eligible: BPL-confirmed current"] == "0"
    # SRC-13 was activated on 2026-10-25 (DEC-15): it moved from Blocked to Authority pending, so nothing is blocked now.
    # It is still not BPL-confirmed, so "Eligible" stays 0.
    assert c["Authority pending"] == "29"
    assert c["Blocked (parse check failed or extraction blocked)"] == "0"
    assert c["Excluded from the corpus"] == "1"
    assert c["Superseded"] == "0"
    assert c["those five lines add up to received"] == "30"
    assert c["Marked as test-index members by the workbook"] == "8"
    assert c["BPL documents (not evaluation material)"] == "29"
    assert rep.disagreements == []  # the shipped workbook agrees with itself


def test_evaluated_is_never_given_a_number() -> None:
    rep = cr.build_report(cr.read_sources(INVENTORY), cr.CorpusFacts())
    line = next(c for c in rep.counts if c.label == "Evaluated")
    assert line.value == "not tracked" and "ING-9" in line.note


def test_unique_by_content_is_not_claimed_without_the_files() -> None:
    rep = cr.build_report(cr.read_sources(INVENTORY), cr.CorpusFacts())
    assert counts_of(rep)["Unique by content"] == "not checked here"


def test_unreviewed_eligible_row_is_not_counted_as_eligible(tmp_path: Path) -> None:
    rows = sources(tmp_path, [wrow("S1", "a.pdf", life="Eligible"), wrow("S2", "b.pdf", life="Eligible", reviewer="B. Bilissalou / 2026-10-12")])
    rep = cr.build_report(rows, cr.CorpusFacts())
    assert rep.eligible == 1  # only the row with a reviewer and a date
    assert counts_of(rep)["Eligible: BPL-confirmed current"] == "1"
    assert counts_of(rep)["Authority pending"] == "1"  # the unreviewed one is pending, as the system treats it
    assert "row_warning" in codes(rep, "S1")


def test_superseded_and_blocked_and_excluded_are_separate_buckets(tmp_path: Path) -> None:
    rows = sources(tmp_path, [
        wrow("S1", "a.pdf"), wrow("S2", "b.pdf", life="Superseded by S1"), wrow("S3", "c.pdf", parse="FAIL: scanned"),
        wrow("S4", "d.docx", life="Excluded", typ="Evaluation material"),
    ])
    c = counts_of(cr.build_report(rows, cr.CorpusFacts()))
    assert (c["Authority pending"], c["Superseded"], c["Blocked (parse check failed or extraction blocked)"], c["Excluded from the corpus"]) == ("1", "1", "1", "1")
    assert c["those five lines add up to received"] == c["Received (rows in the inventory)"] == "4"
    assert c["BPL documents (not evaluation material)"] == "3"


# ------------------------------------------------------------------------------------------------ workbook Counts sheet
GOOD_COUNTS = {"Files received": 2, "BPL documents (excluding test script)": 2, "Eligible (BPL-confirmed current)": 0,
               "Authority pending": 2, "Extraction blocked (parse check failed)": 0, "Excluded from corpus": 0,
               "Currently in the pilot index (not confirmed eligible)": 1, "Indexed AND eligible": 0,
               "Documents with a named open decision (OD-02/03/13)": 0}


def test_counts_sheet_that_agrees_raises_nothing(tmp_path: Path) -> None:
    rows = sources(tmp_path, [wrow("S1", "a.pdf", pilot="Yes"), wrow("S2", "b.pdf")], GOOD_COUNTS)
    rep = cr.build_report(rows, cr.CorpusFacts(), counts_sheet=cr.read_counts_sheet(tmp_path / "inv.xlsx"))
    assert rep.disagreements == []


@pytest.mark.parametrize("label,wrong", [("Files received", 3), ("Authority pending", 1), ("Currently in the pilot index (not confirmed eligible)", 2)])
def test_a_stale_counts_sheet_is_flagged_with_both_numbers(tmp_path: Path, label: str, wrong: int) -> None:
    rows = sources(tmp_path, [wrow("S1", "a.pdf", pilot="Yes"), wrow("S2", "b.pdf")], {**GOOD_COUNTS, label: wrong})
    rep = cr.build_report(rows, cr.CorpusFacts(), counts_sheet=cr.read_counts_sheet(tmp_path / "inv.xlsx"))
    hit = [f for f in rep.disagreements if f.code == "counts_sheet_differs" and f.subject == label]
    assert len(hit) == 1 and f"says {wrong}" in hit[0].text and "recalculate" in hit[0].todo


def test_counts_sheet_with_no_stored_value_is_a_note_not_a_disagreement(tmp_path: Path) -> None:
    rows = sources(tmp_path, [wrow("S1", "a.pdf")], {**GOOD_COUNTS, "Files received": None})
    rep = cr.build_report(rows, cr.CorpusFacts(), counts_sheet=cr.read_counts_sheet(tmp_path / "inv.xlsx"))
    assert any(f.code == "counts_sheet_empty" and f.severity == cr.NOTE for f in rep.findings)


def test_missing_counts_sheet_is_tolerated(tmp_path: Path) -> None:
    make_workbook(tmp_path / "inv.xlsx", [wrow("S1", "a.pdf")])
    assert cr.read_counts_sheet(tmp_path / "inv.xlsx") is None


# ------------------------------------------------------------------------------------------------ never imply readiness
def test_without_a_master_list_no_denominator_is_ever_printed_and_the_gap_is_stated() -> None:
    rep = cr.build_report(cr.read_sources(INVENTORY), cr.CorpusFacts())
    text = cr.render(rep)
    assert "Uncovered policy areas: unknown" in text
    assert "not coverage of the full policy set" in text
    assert not re.search(r"\b8[08]\b", text), "a figure near 80 to 88 reached the output"
    assert not re.search(r"\bof \d+ (policies|documents|sources)\b", text)


def test_headline_always_states_the_confirmed_count_and_calls_test_index_documents_unconfirmed(tmp_path: Path) -> None:
    rows = sources(tmp_path, [wrow("S1", "a.pdf", pilot="Yes"), wrow("S2", "b.pdf", pilot="Yes")])
    both = cr.build_report(rows, facts_with(doc("a.pdf"), doc("b.pdf")))
    assert "2 documents can answer" in both.headline
    assert "All 2 are test-index documents that BPL has not confirmed as current" in both.headline
    assert "BPL-confirmed (eligible) documents: 0" in both.headline
    none = cr.build_report(rows, facts_with())
    assert none.headline.startswith("No document can answer questions now")
    assert "BPL-confirmed (eligible) documents: 0" in none.headline
    skipped = cr.build_report(rows, cr.CorpusFacts())
    assert "cannot say what can answer now" in skipped.headline and "2 as test-index members" in skipped.headline
    assert "BPL-confirmed (eligible) documents: 0" in skipped.headline


# ------------------------------------------------------------------------------------------------ title matching, master list
@pytest.mark.parametrize("text,is_file,expected", [
    ("02 - Materials Selection Policy (04.05) (1) (1).pdf", True, "materials selection policy"),
    ("Materials Selection Policy", False, "materials selection policy"),
    ("02a - Emergency System-wide Closing Procedure (STAFF) (10.02).pdf", True, "emergency system wide closing procedure"),
    ("sub/dir/01 - Technology Use (6.13).pdf", True, "technology use"),
    ("Standards Rev.2026", False, "standards rev 2026"),
])
def test_policy_title_key(text: str, is_file: bool, expected: str) -> None:
    assert cr.policy_title_key(text, is_file=is_file) == expected


def test_near_miss_titles_do_not_match() -> None:
    assert cr.policy_title_key("Technology Use Policy") != cr.policy_title_key("01 - Technology Use (6.13).pdf", is_file=True)
    assert cr.policy_title_key("Technology Use") != cr.policy_title_key("12 - Technology Usage Policy (06.06).pdf", is_file=True)


def test_master_list_reading(tmp_path: Path) -> None:
    (tmp_path / "m.csv").write_text("\ufefftitle,source_id,area\nAnnual Leave Policy,,HR\nDress Code,SRC-17,HR\n,,\n", encoding="utf-8")
    got = cr.read_master_list(tmp_path / "m.csv")
    assert [(e.title, e.source_id, e.area) for e in got] == [("Annual Leave Policy", None, "HR"), ("Dress Code", "SRC-17", "HR")]
    (tmp_path / "m.txt").write_text("# comment\nA\n\n B \n", encoding="utf-8")
    assert [e.title for e in cr.read_master_list(tmp_path / "m.txt")] == ["A", "B"]
    (tmp_path / "bad.csv").write_text("name\nA\n", encoding="utf-8")
    with pytest.raises(ValueError, match="'title'"):
        cr.read_master_list(tmp_path / "bad.csv")
    (tmp_path / "x.xlsx").write_text("x")
    with pytest.raises(ValueError, match=".csv"):
        cr.read_master_list(tmp_path / "x.xlsx")
    (tmp_path / "empty.txt").write_text("# nothing\n")
    with pytest.raises(ValueError, match="no titles"):
        cr.read_master_list(tmp_path / "empty.txt")


def test_coverage_statuses_denominator_and_wording(tmp_path: Path) -> None:
    rows = sources(tmp_path, [
        wrow("S1", "01 - Alpha Policy (1.1).pdf", pilot="Yes"),
        wrow("S2", "02 - Beta Policy (2.2).pdf"),
        wrow("S3", "03 - Gamma Policy.pdf", parse="FAIL: scanned"),
        wrow("S4", "04 - Twin (1).pdf"), wrow("S5", "05 - Twin (2).pdf"),
        wrow("S6", "06 - Unlisted Extra.pdf"),
    ])
    facts = facts_with(doc("01 - Alpha Policy (1.1).pdf"), doc("02 - Beta Policy (2.2).pdf", member=False, answerable=False))
    master = [
        cr.MasterEntry("Alpha Policy", None, "Operations"),
        cr.MasterEntry("Beta", "S2", "HR"),       # linked by source id
        cr.MasterEntry("Gamma Policy", None, "HR"),
        cr.MasterEntry("Twin", None, "HR"),       # two files share this title: refuse to guess
        cr.MasterEntry("Delta Policy", None, "Operations"),   # nothing received
        cr.MasterEntry("Alphas Policy", None, "Operations"),  # near miss must NOT match Alpha
        cr.MasterEntry("Alpha", None, "Operations"),          # a SHORTER title must not match "Alpha Policy"
        cr.MasterEntry("alpha policy!", None, None),          # repeat of the first, differently written
    ]
    rep = cr.build_report(rows, facts, master=master)
    st = {i.entry.title: i.status for i in rep.coverage.items}
    assert st == {"Alpha Policy": "answerable", "Beta": "received_not_answerable", "Gamma Policy": "received_not_answerable",
                  "Twin": "ambiguous", "Delta Policy": "no_file", "Alphas Policy": "no_file", "Alpha": "no_file"}
    assert "master_duplicate" in codes(rep)
    text = cr.render(rep)
    assert "Policies on the master list: 7" in text  # the denominator is the list's own length, repeats counted once
    assert "Have an answerable document now: 1 of 7" in text and "BPL-confirmed documents: 0" in text
    assert "No file matched or ambiguous:    4" in text
    gaps = text.split("No answerable document for:")[1].split("Received but not matched")[0]
    assert "    HR\n" in gaps and "    Operations\n" in gaps and "(no area given)" not in gaps
    assert "- Delta Policy: no received file has this exact title" in gaps and "Alpha Policy" not in gaps
    assert "S6 06 - Unlisted Extra.pdf" in "\n".join(rep.coverage.not_on_list)


def test_a_source_id_that_is_not_in_the_inventory_is_uncovered_not_guessed(tmp_path: Path) -> None:
    rows = sources(tmp_path, [wrow("S1", "a.pdf")])
    rep = cr.build_report(rows, cr.CorpusFacts(), master=[cr.MasterEntry("Whatever", "SRC-99")])
    assert rep.coverage.items[0].status == "no_file" and "not in the inventory" in rep.coverage.items[0].detail


def test_coverage_without_a_database_does_not_claim_answerable(tmp_path: Path) -> None:
    rows = sources(tmp_path, [wrow("S1", "01 - Alpha Policy.pdf", pilot="Yes")])
    rep = cr.build_report(rows, cr.CorpusFacts(), master=[cr.MasterEntry("Alpha Policy")])
    assert rep.coverage.items[0].status == "received_unchecked"
    assert "answerable is unknown" in cr.render(rep) and "Have an answerable document" not in cr.render(rep)


# ------------------------------------------------------------------------------------------------ files on disk
def test_scan_files_hashes_skips_litter_and_finds_duplicates(tmp_path: Path) -> None:
    root = tmp_path / "in"
    (root / "sub").mkdir(parents=True)
    (root / "__MACOSX").mkdir()
    (root / "a.pdf").write_bytes(b"AAA")
    (root / "sub" / "copy of a.pdf").write_bytes(b"AAA")
    (root / "b.pdf").write_bytes(b"BBB")
    (root / ".DS_Store").write_bytes(b"x")
    (root / "__MACOSX" / "._a.pdf").write_bytes(b"x")
    (root / "notes.xyz").write_bytes(b"x")
    scan = cr.scan_files(root)
    assert sorted(e.rel for e in scan.entries) == ["a.pdf", "b.pdf", "sub/copy of a.pdf"]
    assert scan.unique == 2 and scan.ignored == 1
    assert scan.duplicate_groups() == [["a.pdf", "sub/copy of a.pdf"]]


def test_file_findings_against_the_inventory(tmp_path: Path) -> None:
    root = tmp_path / "in"
    root.mkdir()
    (root / "a.pdf").write_bytes(b"A")
    (root / "b (1).pdf").write_bytes(b"B")        # copy marker is ignored when matching, as in apply-inventory
    (root / "stray.pdf").write_bytes(b"C")
    (root / "twin.pdf").write_bytes(b"A")
    rows = sources(tmp_path, [wrow("S1", "a.pdf"), wrow("S2", "b.pdf"), wrow("S3", "gone.pdf")])
    rep = cr.build_report(rows, cr.CorpusFacts(), files=cr.scan_files(root))
    assert {f.subject for f in rep.findings if f.code == "file_without_row"} == {"stray.pdf", "twin.pdf"}
    assert any(f.code == "duplicate_content" and "a.pdf" in f.text and "twin.pdf" in f.text for f in rep.disagreements)
    assert [f.subject for f in rep.findings if f.code == "row_without_file"] == ["S3"]
    assert next(f for f in rep.findings if f.code == "row_without_file").severity == cr.NOTE
    assert counts_of(rep)["Unique by content"] == "3 of 4 files hashed"


def test_the_real_policies_folder_matches_the_real_inventory_if_present() -> None:
    folder = REPO.parent / "Policies"
    if not folder.is_dir():
        pytest.skip("the 30-file Policies folder is not beside the repo")
    rep = cr.build_report(cr.read_sources(INVENTORY), cr.CorpusFacts(), files=cr.scan_files(folder))
    assert not [f for f in rep.findings if f.code in ("file_without_row", "duplicate_content")]


# ------------------------------------------------------------------------------------------------ database (real Postgres, scoped)
def _pinned() -> Settings:
    return get_settings()


async def _make_docs(specs: list[dict[str, Any]]) -> list[uuid.UUID]:
    ids: list[uuid.UUID] = []
    async with async_session_factory() as s:
        for sp in specs:
            did = sp.get("id") or uuid.uuid4()
            s.add(Document(
                document_id=did, title="t", version="t", policy_type="generic_policy", source_format=SourceFormat.md,
                ingest_status=sp.get("status", IngestStatus.indexed), eligibility=sp.get("elig", Eligibility.authority_pending),
                test_index_member=sp.get("member", True), source_path=sp["path"], source_id=sp.get("sid"),
                superseded_by_document_id=sp.get("sup"), extraction_status=sp.get("extraction", ExtractionStatus.standard),
            ))
            ids.append(did)
            await s.flush()
        for sp, did in zip(specs, ids):
            for i in range(sp.get("chunks", 1)):
                s.add(Chunk(chunk_id=uuid.uuid4(), document_id=did, ordinal=i, text="Staff must record gifts.", token_count=4))
        await s.commit()
    return ids


async def _cleanup(ids: list[uuid.UUID]) -> None:
    async with async_session_factory() as s:
        await s.execute(delete(Chunk).where(Chunk.document_id.in_(ids)))
        await s.execute(delete(Document).where(Document.document_id.in_(ids)))
        await s.commit()


def _gather(ids: list[uuid.UUID], settings: Optional[Settings] = None) -> cr.CorpusFacts:
    return asyncio.run(cr.gather_corpus_facts(settings or _pinned(), cr.alembic_head(), scope=ids))


def test_every_disagreement_is_found_against_a_real_database(tmp_path: Path) -> None:
    t = uuid.uuid4().hex[:8]
    n = lambda k: f"zz{t}-{k}.pdf"  # noqa: E731 - unique names so nothing collides with real documents
    sid = lambda k: f"T{t}-{k}"  # noqa: E731
    rows = sources(tmp_path, [
        wrow(sid("ok"), n("ok"), pilot="Yes"),
        wrow(sid("blk"), n("blk"), life="Excluded", typ="Evaluation material"),
        wrow(sid("elg"), n("elg")),
        wrow(sid("mem"), n("mem"), pilot="Yes"),
        wrow(sid("none"), n("none"), pilot="Yes"),
        wrow(sid("sup"), n("sup")),
        wrow(sid("fail"), n("fail"), pilot="Yes"),
        wrow(sid("nochunk"), n("nochunk"), pilot="Yes"),
        wrow(sid("rev"), n("rev"), pilot="Yes"),
        wrow(sid("dup"), n("dup"), pilot="Yes"),
        wrow(sid("sid"), n("sid"), pilot="Yes"),
    ])
    successor = uuid.uuid4()
    specs = [
        dict(path=n("ok")),
        dict(path=n("blk")),                                         # database lets it answer; workbook says excluded
        dict(path=n("elg"), elig=Eligibility.eligible, member=False),  # database says eligible; workbook says pending
        dict(path=n("mem"), member=False),                           # test-index flag differs
        dict(path=n("sup"), member=False, sup=successor),            # replaced in the database only
        dict(path=n("fail"), status=IngestStatus.failed),
        dict(path=n("nochunk"), chunks=0),
        dict(path=n("rev"), extraction=ExtractionStatus.review_required),
        dict(path=n("dup")), dict(path=f"pilot/{n('dup')}"),
        dict(path=n("sid"), sid=f"OTHER-{t}"),
        dict(path=f"zz{t}-unlisted.pdf"),
    ]
    specs.insert(0, dict(id=successor, path=f"zz{t}-successor.pdf", member=False))  # must exist before the row that points at it
    ids = asyncio.run(_make_docs(specs))
    try:
        facts = _gather(ids)
        assert facts.checked and len(facts.docs) == len(specs)
        rep = cr.build_report(rows, facts, inventory_name="x")
        by = {s: codes(rep, sid(s)) for s in ("ok", "blk", "elg", "mem", "none", "sup", "fail", "nochunk", "rev", "dup", "sid")}
        assert by["ok"] == set()
        assert by["blk"] == {"answerable_but_not_allowed", "test_index_differs"}
        assert by["elg"] == {"eligibility_differs"}
        assert by["mem"] == {"test_index_differs"}
        assert by["none"] == {"not_ingested"}
        assert by["sup"] == {"superseded_in_database_only"}
        assert by["fail"] == {"not_indexed"}
        assert by["nochunk"] == {"no_chunks"}
        assert by["rev"] == {"extraction_review_required"}
        assert by["dup"] == {"duplicate_documents"}
        assert by["sid"] == {"source_id_differs"}
        assert any(f.code == "unlisted_document" and f.subject == f"zz{t}-unlisted.pdf" for f in rep.disagreements)
        stage = {ln.source_id: ln.stage for ln in rep.lines}
        assert stage[sid("ok")] == "answerable (test index)"
        assert stage[sid("elg")] == "answerable (BPL-confirmed)"   # the database's truth, which the workbook does not share
        assert stage[sid("none")] == "not ingested" and stage[sid("fail")] == "stored, failed"
        assert stage[sid("nochunk")] == "indexed, no chunks" and stage[sid("blk")] == "answerable (test index)"
        # a stored document with no text is not counted as answerable, whatever the retrieval rule says
        assert next(d for d in facts.docs if d.source_path == n("nochunk")).answerable is False
    finally:
        asyncio.run(_cleanup(ids))


def test_the_review_required_document_cannot_answer_even_though_the_workbook_expects_it(tmp_path: Path) -> None:
    t = uuid.uuid4().hex[:8]
    rows = sources(tmp_path, [wrow(f"T{t}-r", f"zz{t}-r.pdf", pilot="Yes")])
    ids = asyncio.run(_make_docs([dict(path=f"zz{t}-r.pdf", extraction=ExtractionStatus.review_required)]))
    try:
        facts = _gather(ids)
        assert facts.docs[0].answerable is False
        assert cr.build_report(rows, facts).answerable_now == 0
    finally:
        asyncio.run(_cleanup(ids))


def test_test_index_documents_cannot_answer_when_the_setting_is_off_and_the_report_says_why(tmp_path: Path) -> None:
    t = uuid.uuid4().hex[:8]
    rows = sources(tmp_path, [wrow(f"T{t}-a", f"zz{t}-a.pdf", pilot="Yes")])
    ids = asyncio.run(_make_docs([dict(path=f"zz{t}-a.pdf")]))
    try:
        off = _pinned().model_copy(update={"allow_test_index": False})
        rep = cr.build_report(rows, _gather(ids, off))
        assert rep.answerable_now == 0
        hit = [f for f in rep.disagreements if f.code == "not_answerable_though_expected"]
        assert len(hit) == 1 and "ALLOW_TEST_INDEX is off" in hit[0].text
        assert "ALLOW_TEST_INDEX is off" in cr.render(rep)
    finally:
        asyncio.run(_cleanup(ids))


def test_a_clean_database_gives_no_disagreement_and_the_confirmed_split(tmp_path: Path) -> None:
    t = uuid.uuid4().hex[:8]
    rows = sources(tmp_path, [
        wrow(f"T{t}-1", f"zz{t}-1.pdf", pilot="Yes"),
        wrow(f"T{t}-2", f"zz{t}-2.pdf", life="Eligible", reviewer="B. Bilissalou / 2026-10-12"),
    ])
    ids = asyncio.run(_make_docs([dict(path=f"zz{t}-1.pdf"), dict(path=f"zz{t}-2.pdf", elig=Eligibility.eligible, member=False)]))
    try:
        rep = cr.build_report(rows, _gather(ids))
        assert rep.disagreements == [] and rep.answerable_now == 2 and rep.eligible == 1
        c = counts_of(rep)
        assert c["of which BPL-confirmed (eligible)"] == "1" and c["of which test index only"] == "1"
        assert "1 of them are test-index documents that BPL has not confirmed" in rep.headline
        assert c["Indexed AND eligible"] == "1"
    finally:
        asyncio.run(_cleanup(ids))


def test_stored_document_ids_are_printed_in_full_for_the_supersede_step(tmp_path: Path) -> None:
    t = uuid.uuid4().hex[:8]
    rows = sources(tmp_path, [wrow(f"T{t}-1", f"zz{t}-1.pdf", pilot="Yes")])
    ids = asyncio.run(_make_docs([dict(path=f"zz{t}-1.pdf")]))
    try:
        text = cr.render(cr.build_report(rows, _gather(ids)))
        assert "STORED DOCUMENTS" in text and str(ids[0]) in text
        assert "STORED DOCUMENTS" not in cr.render(cr.build_report(rows, _gather(ids)), summary_only=True)
    finally:
        asyncio.run(_cleanup(ids))


# ------------------------------------------------------------------------------------------------ expected_state mirrors apply
@pytest.mark.parametrize("elig", list(Eligibility))
@pytest.mark.parametrize("member", [False, True])
@pytest.mark.parametrize("claim", [False, True])
@pytest.mark.parametrize("extraction", [ExtractionStatus.standard, ExtractionStatus.validated, ExtractionStatus.review_required])
def test_expected_state_agrees_with_what_apply_inventory_actually_stores(elig, member, claim, extraction) -> None:
    """If these ever drift the report would 'find' disagreements that apply cannot fix, or miss real ones."""
    t = uuid.uuid4().hex[:8]
    name = f"zz{t}-m.pdf"
    row = InventoryRow(source_id=f"T{t}-m", file_name=name, eligibility=elig, test_index_member=member,
                       extraction_validated_claim=claim)
    ids = asyncio.run(_make_docs([dict(path=name, extraction=extraction, member=False)]))
    try:
        before = _gather(ids).docs[0]
        want = cr.expected_state(row, before)

        async def apply() -> None:
            async with async_session_factory() as s:
                await apply_inventory_to_db(s, {row.key: row}, scope=ids)

        asyncio.run(apply())
        after = _gather(ids).docs[0]
        assert (after.eligibility, after.test_index_member) == (want[0].value, want[1])
    finally:
        asyncio.run(_cleanup(ids))


def test_expected_state_with_no_stored_document_and_a_validated_claim_is_blocked() -> None:
    row = InventoryRow(source_id="S", file_name="x.pdf", eligibility=Eligibility.authority_pending, test_index_member=True,
                       extraction_validated_claim=True)
    assert cr.expected_state(row, None) == (Eligibility.blocked, False)


# ------------------------------------------------------------------------------------------------ read-only guarantees
def test_database_connection_is_declared_read_only_and_issues_no_write(monkeypatch: pytest.MonkeyPatch) -> None:
    import sqlalchemy.ext.asyncio as sa_asyncio
    from sqlalchemy import event

    seen: list[str] = []
    real = sa_asyncio.create_async_engine

    def spying(*a: Any, **k: Any):
        eng = real(*a, **k)

        @event.listens_for(eng.sync_engine, "before_cursor_execute")
        def _rec(_c, _cur, statement, *_rest) -> None:  # noqa: ANN001
            seen.append(statement.strip())

        return eng

    monkeypatch.setattr(sa_asyncio, "create_async_engine", spying)
    facts = asyncio.run(cr.gather_corpus_facts(_pinned(), cr.alembic_head(), scope=[uuid.uuid4()]))
    assert facts.checked and seen
    assert seen[0].upper().startswith("SET TRANSACTION READ ONLY")
    for stmt in seen:
        assert not re.match(r"(?is)\s*(insert|update|delete|alter|drop|create|truncate|grant)\b", stmt), stmt


def test_gathering_changes_no_rows() -> None:
    async def counts() -> tuple[int, int]:
        async with async_session_factory() as s:
            return ((await s.execute(select(func.count()).select_from(Document))).scalar_one(),
                    (await s.execute(select(func.count()).select_from(Chunk))).scalar_one())

    before = asyncio.run(counts())
    assert asyncio.run(cr.gather_corpus_facts(_pinned(), cr.alembic_head())).checked
    assert asyncio.run(counts()) == before


def test_source_contains_no_database_write_and_no_network_call() -> None:
    src = (ORCH / "ask_bpl_orchestrator" / "corpus_report.py").read_text()
    code = "\n".join(line for line in src.splitlines() if not line.lstrip().startswith(("#", '"""')))
    assert not re.search(r"\b(INSERT INTO|DELETE FROM|UPDATE \w+ SET|DROP TABLE|ALTER TABLE|TRUNCATE)\b", code)
    assert ".commit(" not in code and "session.add(" not in code and "async_session_factory" not in code
    assert not re.search(r"\b(httpx|requests|urllib|pinecone|openai|anthropic)\b", code), "the report must never call an outside service"


def test_the_only_file_written_is_the_csv_and_only_when_asked(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys) -> None:
    monkeypatch.chdir(tmp_path)
    before = sorted(p.name for p in tmp_path.iterdir())
    assert cr.main(["--no-db", "--inventory", str(INVENTORY)]) == 0
    assert sorted(p.name for p in tmp_path.iterdir()) == before
    out = tmp_path / "t.csv"
    assert cr.main(["--no-db", "--inventory", str(INVENTORY), "--csv", str(out)]) == 0
    assert [p.name for p in tmp_path.iterdir() if p.name not in before] == ["t.csv"]
    rows = list(csv.reader(out.open(encoding="utf-8")))
    assert rows[0][:3] == ["source_id", "file_name", "source_type"] and len(rows) == 31
    capsys.readouterr()


def test_csv_cells_cannot_start_a_spreadsheet_formula(tmp_path: Path) -> None:
    rows = sources(tmp_path, [wrow("S1", "+cmd.pdf")])
    cr.write_csv(cr.build_report(rows, cr.CorpusFacts()), tmp_path / "o.csv")
    assert list(csv.reader((tmp_path / "o.csv").open()))[1][1] == "'+cmd.pdf"


# ------------------------------------------------------------------------------------------------ failure to read the database
def test_unreachable_database_is_reported_in_fixed_words_and_never_prints_the_password() -> None:
    pw = "ZZTOPdbpass9999"
    s = Settings(_env_file=None, database_url=f"postgresql+asyncpg://askbpl:{pw}@127.0.0.1:1/none", jwt_secret_key="x" * 40,
                 openai_api_key="a", anthropic_api_key="b", pinecone_api_key="c", pinecone_index_name="i")
    facts = asyncio.run(cr.gather_corpus_facts(s, None))
    assert facts.state == "unreachable" and facts.error_kind in ("cannot_connect", "timeout", "unknown") and facts.docs == []
    rep = cr.build_report(cr.read_sources(INVENTORY), facts)
    text = cr.render(rep)
    assert pw not in text and "ZZTOP" not in text and "askbpl:" not in text
    assert "could not be reached" in text and "cannot say what can answer now" in text
    assert not rep.disagreements  # an unreadable database is not a finding about the documents


def test_out_of_date_database_is_reported_not_queried() -> None:
    facts = asyncio.run(cr.gather_corpus_facts(_pinned(), "not-the-real-head"))
    assert facts.state == "out_of_date" and facts.docs == []
    assert "older than this code" in cr.render(cr.build_report(cr.read_sources(INVENTORY), facts))


# ------------------------------------------------------------------------------------------------ command line
def test_exit_codes(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys) -> None:
    monkeypatch.chdir(tmp_path)
    assert cr.main(["--inventory", "nope.xlsx", "--no-db"]) == 2
    (tmp_path / "bad.csv").write_text("name\nx\n")
    assert cr.main(["--no-db", "--inventory", str(INVENTORY), "--master-list", str(tmp_path / "bad.csv")]) == 2
    assert cr.main(["--no-db", "--inventory", str(INVENTORY), "--files", str(tmp_path / "missing")]) == 2
    bad_life = make_workbook(tmp_path / "life.xlsx", [wrow("S1", "a.pdf", life="Approved-ish")])
    assert cr.main(["--no-db", "--inventory", str(bad_life)]) == 2  # the system's fail-closed rule, not a guess
    assert cr.main(["--no-db", "--inventory", str(INVENTORY), "--strict"]) == 0
    stale = make_workbook(tmp_path / "stale.xlsx", [wrow("S1", "a.pdf")], {**GOOD_COUNTS, "Files received": 9})
    assert cr.main(["--no-db", "--inventory", str(stale)]) == 0          # findings are not errors...
    assert cr.main(["--no-db", "--inventory", str(stale), "--strict"]) == 1  # ...unless --strict
    capsys.readouterr()


def test_strict_fails_when_the_database_was_asked_for_but_cannot_be_read(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys) -> None:
    async def down(*_a: Any, **_k: Any) -> cr.CorpusFacts:
        return cr.CorpusFacts(state="unreachable", error_kind="cannot_connect")

    monkeypatch.setattr(cr, "gather_corpus_facts", down)
    assert cr.main(["--inventory", str(INVENTORY)]) == 0
    assert cr.main(["--inventory", str(INVENTORY), "--strict"]) == 1
    assert cr.main(["--inventory", str(INVENTORY), "--strict", "--no-db"]) == 0
    capsys.readouterr()


def test_the_default_inventory_is_found_from_any_folder(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys) -> None:
    monkeypatch.chdir(tmp_path)
    assert cr.main(["--no-db", "--summary-only"]) == 0
    assert "Received (rows in the inventory)" in capsys.readouterr().out


def test_command_is_registered() -> None:
    assert 'askbpl-corpus-report = "ask_bpl_orchestrator.corpus_report:main"' in (ORCH / "pyproject.toml").read_text()


# ------------------------------------------------------------------------------------------------ the curator guide
GUIDE = REPO / "docs" / "CURATOR_WORKFLOW.md"


def _scripts() -> dict[str, str]:
    return dict(re.findall(r'^(askbpl-[a-z\-]+) = "([^"]+)"', (ORCH / "pyproject.toml").read_text(), re.M))


def _logical_commands(guide: str) -> list[str]:
    out: list[str] = []
    for block in re.findall(r"```[a-z]*\n(.*?)```", guide, re.S):
        for line in re.sub(r"\\\n\s*", " ", block).splitlines():
            line = line.split("   #")[0].strip()
            if line.startswith("askbpl-"):
                out.append(line)
    return out


def test_guide_exists_and_names_only_real_commands() -> None:
    guide = GUIDE.read_text()
    named = {c.split()[0] for c in _logical_commands(guide)}
    assert {"askbpl-doctor", "askbpl-corpus-report", "askbpl-apply-inventory", "askbpl-ingest-batch", "askbpl-ingest",
            "askbpl-extraction-review", "askbpl-eval", "askbpl-reader-compare", "askbpl-provenance-report"} <= named
    assert named <= set(_scripts()), named - set(_scripts())
    # inline mentions too (anything that looks like a command in backticks)
    inline = set(re.findall(r"`(askbpl-[a-z]+(?:-[a-z]+)*)[ `]", guide))
    assert inline <= set(_scripts()), inline - set(_scripts())


def test_every_flag_the_guide_uses_exists_in_that_commands_help(monkeypatch: pytest.MonkeyPatch, capsys) -> None:
    import importlib

    checked = 0
    for line in _logical_commands(GUIDE.read_text()):
        tokens = shlex.split(line)
        name, rest = tokens[0], tokens[1:]
        flags = [t for t in rest if t.startswith("--")]
        if not flags:
            continue
        sub = [rest[0]] if rest and not rest[0].startswith("-") and name == "askbpl-extraction-review" else []
        module, func = _scripts()[name].split(":")
        monkeypatch.setattr(sys, "argv", [name, *sub, "--help"])
        with pytest.raises(SystemExit):
            getattr(importlib.import_module(module), func)()
        helptext = capsys.readouterr().out
        for f in flags:
            assert f in helptext, f"{name} {' '.join(sub)}: the guide uses {f}, which --help does not list"
            checked += 1
    assert checked >= 12


def test_guide_covers_every_curator_action_and_says_what_is_undecided() -> None:
    g = GUIDE.read_text()
    for heading in ("### A. Add a new policy", "### B. Replace a file in place", "### C. Supersede", "### D. Retire a policy",
                    "### E. Re-ingest after a setting change", "### F. A file was renamed or moved", "## 5. Validate"):
        assert heading in g, heading
    low = " ".join(g.lower().split())  # line wraps must not hide a phrase
    for needle in ("open bpl decisions", "marked test index", "outside services", "not yet validated with bpl", "od-01",
                   "no command that deletes", "not recorded", "cannot see pinecone", "reviewer **and** a review date"):
        assert needle in low, needle
    assert not re.search(r"sk-[A-Za-z0-9_\-]{8,}|pcsk_[A-Za-z0-9]{8,}", g)


def test_guide_never_says_a_document_is_approved_or_that_88_are_ready() -> None:
    g = GUIDE.read_text()
    assert not re.search(r"\b(88|80) (policies|documents|sources) (are|is) ready\b", g, re.I)
    assert "nothing is bpl-confirmed" in g.lower()


# --- pins for the behaviours the guide warns about. If one fails because the code was improved, fix the guide too.
def test_guide_claim_document_identity_is_the_relative_path_ignoring_case() -> None:
    from ask_bpl_orchestrator.ingest.service import _document_uuid

    assert _document_uuid("Pilot/A Policy.pdf") == _document_uuid("pilot/a policy.pdf")
    assert _document_uuid("pilot/a policy.pdf") != _document_uuid("policies/a policy.pdf")  # a move is a different document


def test_guide_claim_a_superseded_document_is_not_opened_by_the_citation_viewer() -> None:
    from ask_bpl_orchestrator.services.eligibility import is_retrievable

    d = SimpleNamespace(ingest_status=IngestStatus.indexed, superseded_by_document_id=uuid.uuid4(), extraction_status="standard",
                        eligibility=Eligibility.eligible, test_index_member=False)
    assert is_retrievable(d) is False
    route = (ORCH / "ask_bpl_orchestrator" / "api" / "routes_documents.py").read_text()
    assert "not is_retrievable(doc)" in route and "404" in route.replace("HTTP_404_NOT_FOUND", "404")


@pytest.mark.asyncio
async def test_batch_ingest_refuses_blocked_excluded_and_superseded_rows(tmp_path: Path) -> None:
    """Patch 0021 closed the gap where a Superseded row was uploaded again; the curator guide says so."""
    from ask_bpl_orchestrator.ingest import batch

    corpus = tmp_path / "corpus"
    corpus.mkdir()
    for f in ("old.md", "blocked.md", "new.md"):
        (corpus / f).write_text("# x\n\nStaff must record gifts.\n")
    (tmp_path / "m.yaml").write_text(
        "corpus_root: corpus\ndocuments:\n  - {path: old.md, policy: generic_policy, title: Old}\n"
        "  - {path: blocked.md, policy: generic_policy, title: B}\n  - {path: new.md, policy: generic_policy, title: New}\n")
    inv = make_workbook(tmp_path / "i.xlsx", [wrow("S1", "old.md", life="Superseded by S3"),
                                              wrow("S2", "blocked.md", parse="FAIL: scanned"), wrow("S3", "new.md")])
    called: list[str] = []

    async def fake(**kw: Any) -> None:
        called.append(kw["absolute_path"].name)

    with patch.object(batch, "ingest_file", side_effect=fake):
        lines = await batch.ingest_manifest(tmp_path / "m.yaml", inv)
    assert called == ["new.md"]
    assert any(ln.startswith("refused:old.md:superseded:S1") for ln in lines)
    assert any(ln.startswith("refused:blocked.md:blocked:S2") for ln in lines)

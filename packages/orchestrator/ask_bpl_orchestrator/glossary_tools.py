"""Read-only glossary tools (patch 0016; PRD GLO-3, GLO-5).

    askbpl-glossary-check [path/to/glossary.yaml]
    askbpl-glossary-candidates --general "General.xlsx" --policy "Policy terms.xlsx" [--pdf P.pdf ...]
                               [--eval-csv Q.csv ...] [--csv out.csv] [--top 40]

``check`` validates the glossary file the app uses and says exactly what is wrong. ``candidates`` reads the two draft
workbooks (both are drafts, whatever their cells say: GLO-5) and prepares a review list for BPL: which terms occur in
the policies and eval questions, where the two workbooks disagree, which meanings are unknown. It never writes the
glossary file and approves nothing. No database, no keys, no network.
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from ask_bpl_orchestrator.services.glossary import (
    BUNDLED_PATH,
    GlossaryError,
    _is_acronym,
    form_pattern,
    load_glossary,
)

GENERAL_SHEET = "Master Glossary"
POLICY_SHEET = "BPL Glossary"
REVIEW_COLUMNS = [
    "decision",
    "approved_expansion",
    "aliases",
    "owner",
    "approved_by",
    "approved_on",
    "notes",
]


# ----------------------------------------------------------------------------------------------- check
def check_main(argv: Optional[list[str]] = None) -> int:
    parser = argparse.ArgumentParser(description="Validate the glossary file (read-only).")
    parser.add_argument(
        "path", nargs="?", type=Path, default=None, help="default: the bundled config/glossary.yaml"
    )
    args = parser.parse_args(argv)
    path = args.path or BUNDLED_PATH
    try:
        glossary = load_glossary(path)
    except GlossaryError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    counts = Counter(e.status for e in glossary.entries)
    print(f"glossary file: {path}")
    print(
        f"entries: {len(glossary.entries)} (approved {counts['approved']}, draft {counts['draft']}, "
        f"conflict {counts['conflict']}); extra allowlisted words: {len(glossary.allowlist)}"
    )
    if not glossary.approved:
        print("No approved entries: the glossary changes nothing until BPL approves terms.")
    print("OK")
    return 0


# ----------------------------------------------------------------------------------------------- candidates
@dataclass
class Draft:
    term: str
    meaning: str
    category: str
    status: str
    source: str
    workbook: str


@dataclass
class Candidate:
    key: str
    term: str
    general: list[Draft] = field(default_factory=list)
    policy: list[Draft] = field(default_factory=list)
    eval_hits: int = 0
    text_hits: int = 0
    spelled_out_hits: int = 0
    flags: list[str] = field(default_factory=list)

    @property
    def drafts(self) -> list[Draft]:
        return [*self.general, *self.policy]

    @property
    def score(self) -> int:
        return 3 * self.eval_hits + 2 * self.text_hits + self.spelled_out_hits


def _cell(value: object) -> str:
    return "" if value is None else " ".join(str(value).split())


def _norm(meaning: str) -> str:
    text = meaning.casefold().replace("&", " and ")
    text = re.sub(r"^\s*(shorthand|short form|abbreviation) for\s+", "", text)
    return " ".join(re.sub(r"[^a-z0-9 ]+", " ", text).split())


def _rows(path: Path, sheet: str, header_first: str) -> list[dict[str, str]]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet]
    except KeyError as exc:
        raise SystemExit(f"{path.name}: no sheet named {sheet!r} (found {wb.sheetnames})") from exc
    header: Optional[list[str]] = None
    out: list[dict[str, str]] = []
    for row in ws.iter_rows(values_only=True):
        cells = [_cell(c) for c in row]
        if header is None:
            if cells and cells[0] == header_first:
                header = cells
            continue
        if any(cells):
            out.append(dict(zip(header, cells)))
    if header is None:
        raise SystemExit(
            f"{path.name}: header row starting with {header_first!r} not found in {sheet!r}"
        )
    return out


def read_general(path: Path) -> list[Draft]:
    rows = _rows(path, GENERAL_SHEET, "Category")
    return [
        Draft(
            r["Term / Acronym"],
            r.get("Meaning / BPL Context", ""),
            r.get("Category", ""),
            r.get("Verification Status", ""),
            r.get("Source Basis", ""),
            "general",
        )
        for r in rows
        if r.get("Term / Acronym")
    ]


def read_policy(path: Path) -> list[Draft]:
    rows = _rows(path, POLICY_SHEET, "Acronym / Term")
    return [
        Draft(
            r["Acronym / Term"],
            r.get("Meaning / Expansion", ""),
            r.get("Category", ""),
            r.get("Verification Status", ""),
            r.get("Source File(s)", ""),
            "policy",
        )
        for r in rows
        if r.get("Acronym / Term")
    ]


def merge(general: list[Draft], policy: list[Draft]) -> list[Candidate]:
    by_key: dict[str, Candidate] = {}
    for draft in [*general, *policy]:
        key = draft.term.casefold().strip()
        cand = by_key.setdefault(key, Candidate(key=key, term=draft.term))
        (cand.general if draft.workbook == "general" else cand.policy).append(draft)
    for cand in by_key.values():
        cand.flags = _flags(cand)
    return list(by_key.values())


def _flags(cand: Candidate) -> list[str]:
    flags: list[str] = []
    if cand.general and cand.policy:
        flags.append("in both workbooks")
        a, b = _norm(cand.general[0].meaning), _norm(cand.policy[0].meaning)
        if a and b and a not in b and b not in a:
            flags.append("MEANINGS DIFFER")
    if any(not d.meaning or "not stated" in d.status.casefold() for d in cand.drafts):
        flags.append("unknown meaning")
    if any(d.workbook == "general" and not d.source for d in cand.drafts):
        flags.append("no source")
    if len(cand.general) > 1 or len(cand.policy) > 1:
        flags.append("duplicate in a workbook")
    return flags


def _expansion_key(cand: Candidate) -> str:
    """The long form to look for in policy text, only when it is short enough to be a name and not a description."""
    for draft in cand.drafts:
        meaning = _norm(draft.meaning)
        if meaning and len(meaning.split()) <= 6:
            return meaning
    return ""


def count_hits(candidates: list[Candidate], questions: list[str], policy_texts: list[str]) -> None:
    blob = "\n".join(policy_texts)
    folded = " ".join(blob.casefold().split())
    for cand in candidates:
        pattern = form_pattern(cand.term)
        cand.eval_hits = sum(1 for q in questions if pattern.search(q))
        cand.text_hits = len(pattern.findall(blob))
        expansion = _expansion_key(cand)
        cand.spelled_out_hits = (
            folded.count(expansion) if expansion and expansion != cand.key else 0
        )


def _questions(paths: list[Path]) -> list[str]:
    out: list[str] = []
    for path in paths:
        with open(path, newline="", encoding="utf-8") as fh:
            out.extend(row["question"] for row in csv.DictReader(fh) if row.get("question"))
    return out


def _policy_texts(paths: list[Path]) -> list[str]:
    from ask_bpl_orchestrator.ingest.loaders import load_bytes

    return [load_bytes(p) for p in paths]


def _meaning(cand: Candidate, which: str) -> str:
    drafts = cand.general if which == "general" else cand.policy
    return drafts[0].meaning if drafts else ""


def report(
    candidates: list[Candidate],
    general: list[Draft],
    policy: list[Draft],
    *,
    n_pdfs: int,
    n_questions: int,
    top: int,
    acronyms_only: bool = False,
) -> str:
    g_terms = {d.term.casefold() for d in general}
    lines = [
        "=== GLOSSARY CANDIDATES (both workbooks are drafts; this tool approves nothing and writes no glossary file)",
        f"general workbook: {len(general)} rows, {len(g_terms)} unique terms; statuses {dict(Counter(d.status for d in general))}",
        f"policy workbook: {len(policy)} rows; statuses {dict(Counter(d.status for d in policy))}",
        f"terms in both: {sum('in both workbooks' in c.flags for c in candidates)}   "
        f"meanings differ: {sum('MEANINGS DIFFER' in c.flags for c in candidates)}   "
        f"unknown meaning: {sum('unknown meaning' in c.flags for c in candidates)}   "
        f"no source: {sum('no source' in c.flags for c in candidates)}   "
        f"duplicate in a workbook: {sum('duplicate in a workbook' in c.flags for c in candidates)}",
        f"text searched: {n_pdfs} policy PDF(s), {n_questions} eval question(s)",
    ]
    ranked = sorted(
        (c for c in candidates if c.score > 0 and (not acronyms_only or _is_acronym(c.term))),
        key=lambda c: (-c.score, c.key),
    )
    kind = "acronyms" if acronyms_only else "terms"
    lines.append(
        f"\n--- Prioritised: {kind} that appear in the searched text ({len(ranked)} of {len(candidates)} entries), top {top}"
    )
    if not ranked:
        lines.append("  none (pass --pdf and --eval-csv to rank terms by use)")
    for c in ranked[:top]:
        meaning = _meaning(c, "general") or _meaning(c, "policy")
        lines.append(
            f"  {c.term:<14} score {c.score:<3} eval {c.eval_hits} | in policy text {c.text_hits} | spelled out {c.spelled_out_hits}"
            f" | {', '.join(c.flags) or '-'} | {meaning[:70]}"
        )
    differ = sorted((c for c in candidates if "MEANINGS DIFFER" in c.flags), key=lambda c: c.key)
    lines.append(f"\n--- Review first: the two workbooks disagree ({len(differ)})")
    for c in differ:
        lines.append(
            f"  {c.term}: general = {_meaning(c, 'general')[:60]!r} | policy = {_meaning(c, 'policy')[:60]!r}"
        )
    return "\n".join(lines)


def write_csv(candidates: list[Candidate], path: Path) -> None:
    columns = [
        "term",
        "category",
        "meaning_general",
        "meaning_policy",
        "status_general",
        "status_policy",
        "flags",
        "eval_hits",
        "policy_text_hits",
        "spelled_out_hits",
        *REVIEW_COLUMNS,
    ]
    with open(path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(columns)
        for c in sorted(candidates, key=lambda c: (-c.score, c.key)):
            drafts = c.drafts
            writer.writerow(
                [
                    c.term,
                    drafts[0].category if drafts else "",
                    _meaning(c, "general"),
                    _meaning(c, "policy"),
                    c.general[0].status if c.general else "",
                    c.policy[0].status if c.policy else "",
                    "; ".join(c.flags),
                    c.eval_hits,
                    c.text_hits,
                    c.spelled_out_hits,
                    *([""] * len(REVIEW_COLUMNS)),
                ]
            )


def candidates_main(argv: Optional[list[str]] = None) -> int:
    parser = argparse.ArgumentParser(
        description="Prepare a glossary review list from the two draft workbooks (read-only)."
    )
    parser.add_argument(
        "--general",
        type=Path,
        required=True,
        help="the general terms workbook (sheet 'Master Glossary')",
    )
    parser.add_argument(
        "--policy",
        type=Path,
        required=True,
        help="the policy terms workbook (sheet 'BPL Glossary')",
    )
    parser.add_argument(
        "--pdf", type=Path, nargs="*", default=[], help="policy PDFs to count term use in"
    )
    parser.add_argument(
        "--eval-csv",
        type=Path,
        nargs="*",
        default=[],
        help="eval case files (the 'question' column)",
    )
    parser.add_argument(
        "--csv",
        type=Path,
        default=None,
        help="write every candidate, with blank review columns, here",
    )
    parser.add_argument("--top", type=int, default=40)
    parser.add_argument(
        "--acronyms-only",
        action="store_true",
        help="rank only short all-capital terms (BCC, DPS, HR)",
    )
    args = parser.parse_args(argv)
    for path in [args.general, args.policy, *args.pdf, *args.eval_csv]:
        if not path.is_file():
            print(f"not a file: {path}", file=sys.stderr)
            return 2
    general, policy = read_general(args.general), read_policy(args.policy)
    candidates = merge(general, policy)
    questions = _questions(args.eval_csv)
    count_hits(candidates, questions, _policy_texts(args.pdf))
    print(
        report(
            candidates,
            general,
            policy,
            n_pdfs=len(args.pdf),
            n_questions=len(questions),
            top=args.top,
            acronyms_only=args.acronyms_only,
        )
    )
    if args.csv:
        write_csv(candidates, args.csv)
        print(f"\nwritten: {args.csv} ({len(candidates)} rows; review columns left blank)")
    return 0


if __name__ == "__main__":
    raise SystemExit(candidates_main())

"""askbpl-corpus-report: READ-ONLY counts of the policy corpus (patch 0020; PRD ING-13, ING-9, AC-1).

    askbpl-corpus-report                          # inventory workbook + database
    askbpl-corpus-report --no-db                  # workbook only (no database needed)
    askbpl-corpus-report --files <folder>         # also hash the received files (unique by content, files never inventoried)
    askbpl-corpus-report --master-list titles.csv # also say which policies on BPL's list have no answerable document
    askbpl-corpus-report --csv /tmp/corpus.csv    # also write the per-source table (the only thing this command writes)
    askbpl-corpus-report --strict                 # exit 1 when anything needs reconciling (for a pre-demo check)

It answers one question without flattering the team: how many documents have we received, and how many of them can
actually answer a question today? The counts are mutually exclusive buckets that add up to the number received, so
nothing is counted twice (the workbook's own "Authority pending" line also counts the blocked file; this report does
not). It never prints a coverage figure against 80 to 88 policies unless a master list is supplied, because without
that list nobody knows what the denominator is.

It changes nothing: the database connection is declared READ ONLY, and the workbook, the settings file, the glossary,
the policy files and Pinecone are only read. Without ``--csv`` it writes no file at all.

What it cannot tell you (the report says so every time):

* whether a document's vectors in Pinecone match the database (it never calls Pinecone),
* whether answers are right, whether a policy is current, or whether BPL has approved anything,
* how many documents were evaluated: the "Evaluated" stage of ING-9 is not recorded anywhere yet.

Error text from the database is never printed (it can echo a connection string); the fixed sentences from
``doctor.py`` are used instead.
"""

from __future__ import annotations

import argparse
import asyncio
import csv
import hashlib
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence

from ask_bpl_orchestrator.config import Settings
from ask_bpl_orchestrator.db.models import Eligibility, ExtractionStatus, IngestStatus
from ask_bpl_orchestrator.doctor import alembic_head, classify_error, describe_database_target
from ask_bpl_orchestrator.ingest.inventory import InventoryRow, load_inventory, normalize_name

SOURCE_SHEET = "Source Inventory"
COUNTS_SHEET = "Counts (ING-13)"
INGESTIBLE_SUFFIXES = {".pdf", ".docx", ".md", ".markdown", ".txt"}

DISAGREEMENT, NOTE = "disagreement", "note"
STAGE_ANSWERABLE_CONFIRMED = "answerable (BPL-confirmed)"
STAGE_ANSWERABLE_TEST = "answerable (test index)"

NOT_EVALUATED_NOTE = "not tracked: the Evaluated stage of ING-9 is not recorded anywhere yet"
CANNOT_TELL = (
    "Pinecone was not asked, so this report cannot confirm that the stored vectors match the database.",
    "Nothing here shows that answers are right, that a policy is current, or that BPL has approved it.",
    "A document counted as answerable is answerable only because of the test-index setting while BPL authority is pending.",
)


# ------------------------------------------------------------------------------------------------ data
@dataclass
class SourceRow:
    """One inventory row: the system's reading of it (`inv`) plus the raw columns the report prints."""

    inv: InventoryRow
    fmt: str = ""
    parse_check: str = ""
    in_pilot_text: str = ""
    source_type: str = ""
    lifecycle: str = ""
    blocking: str = ""


@dataclass
class DbDoc:
    document_id: str
    title: str
    source_path: Optional[str]
    source_id: Optional[str]
    eligibility: str
    test_index_member: bool
    ingest_status: str
    extraction_status: str
    superseded: bool
    chunks: int
    answerable: bool

    @property
    def key(self) -> str:
        return normalize_name(self.source_path or self.title)


@dataclass
class CorpusFacts:
    """What the database holds. `state` says whether it could be read; the report never raises on a bad database."""

    state: str = "skipped"  # skipped | checked | unreachable | not_migrated | out_of_date
    error_kind: str = ""
    target: str = ""
    allow_test_index: bool = True
    docs: List[DbDoc] = field(default_factory=list)

    @property
    def checked(self) -> bool:
        return self.state == "checked"


@dataclass
class Finding:
    code: str
    severity: str
    subject: str
    text: str
    todo: str = ""


@dataclass
class SourceLine:
    source_id: str
    file_name: str
    source_type: str
    expected: str
    expected_test_index: bool
    stage: str
    doc: Optional[DbDoc]
    finding_codes: List[str] = field(default_factory=list)


@dataclass
class FileEntry:
    rel: str
    sha256: str


@dataclass
class FilesScan:
    root: str
    entries: List[FileEntry] = field(default_factory=list)
    ignored: int = 0
    unreadable: List[str] = field(default_factory=list)

    @property
    def unique(self) -> int:
        return len({e.sha256 for e in self.entries})

    def duplicate_groups(self) -> List[List[str]]:
        by: Dict[str, List[str]] = {}
        for e in self.entries:
            by.setdefault(e.sha256, []).append(e.rel)
        return [sorted(v) for v in by.values() if len(v) > 1]


@dataclass
class MasterEntry:
    title: str
    source_id: Optional[str] = None
    area: Optional[str] = None


@dataclass
class CoverageItem:
    entry: MasterEntry
    status: str  # answerable | received_not_answerable | received_unchecked | no_file | ambiguous
    detail: str = ""


@dataclass
class Coverage:
    items: List[CoverageItem]
    not_on_list: List[str]

    def count(self, status: str) -> int:
        return sum(1 for i in self.items if i.status == status)


@dataclass
class CountLine:
    label: str
    value: str
    note: str = ""


@dataclass
class CorpusReport:
    inventory_name: str
    facts: CorpusFacts
    counts: List[CountLine]
    lines: List[SourceLine]
    findings: List[Finding]
    coverage: Optional[Coverage]
    headline: str
    received: int
    answerable_now: Optional[int]
    eligible: int

    @property
    def disagreements(self) -> List[Finding]:
        return [f for f in self.findings if f.severity == DISAGREEMENT]


# ------------------------------------------------------------------------------------------------ reading the workbook
def _clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def read_sources(path: Path) -> List[SourceRow]:
    """Inventory rows in workbook order. The system's own loader decides eligibility; this adds the raw columns.

    Raises ValueError (the loader's own, fail-closed messages) when the workbook cannot be trusted.
    """
    import openpyxl  # local import: only needed for ingest tooling

    inv = load_inventory(path)  # authority for eligibility, test-index membership, reviewer; raises on bad rows
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        rows = list(wb[SOURCE_SHEET].iter_rows(values_only=True))
    finally:
        wb.close()
    header = [_clean(h) for h in rows[0]]
    col = {name: i for i, name in enumerate(header)}

    def get(r: Sequence[Any], name: str) -> str:
        i = col.get(name)
        return _clean(r[i]) if i is not None and i < len(r) else ""

    out: List[SourceRow] = []
    for r in rows[1:]:
        sid, fname = get(r, "Source ID"), get(r, "File name")
        if not sid or not fname:
            continue
        row = inv[normalize_name(fname)]
        out.append(
            SourceRow(
                inv=row,
                fmt=get(r, "Format"),
                parse_check=get(r, "Parse check"),
                in_pilot_text=get(r, "In current pilot index"),
                source_type=get(r, "Source type"),
                lifecycle=get(r, "Lifecycle status (ING-9)"),
                blocking=get(r, "Blocking decision"),
            )
        )
    return out


def read_counts_sheet(path: Path) -> Optional[Dict[str, Any]]:
    """The workbook's own 'Counts (ING-13)' sheet as {measure: stored value}; None if the sheet is absent."""
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        if COUNTS_SHEET not in wb.sheetnames:
            return None
        out: Dict[str, Any] = {}
        for r in wb[COUNTS_SHEET].iter_rows(min_row=2, values_only=True):
            if r and r[0] and len(r) > 1 and not str(r[0]).startswith("Legend"):
                out[str(r[0]).strip()] = r[1]
        return out
    finally:
        wb.close()


def _starts(text: str, prefix: str) -> bool:
    """Excel's COUNTIF("prefix*"): case-insensitive prefix match."""
    return text.strip().lower().startswith(prefix.lower())


def _equals(text: str, value: str) -> bool:
    return text.strip().lower() == value.lower()


def recompute_workbook_counts(rows: Sequence[SourceRow]) -> Dict[str, int]:
    """The measures on the workbook's Counts sheet, recomputed with the same rules its formulas use."""
    od = ("OD-02", "OD-03", "OD-13")
    n_eval = sum(_equals(r.source_type, "Evaluation material") for r in rows)
    return {
        "Files received": len(rows),
        "BPL documents (excluding test script)": len(rows) - n_eval,
        "Eligible (BPL-confirmed current)": sum(_starts(r.lifecycle, "Eligible") for r in rows),
        "Authority pending": sum(_starts(r.lifecycle, "Authority pending") for r in rows),
        "Extraction blocked (parse check failed)": sum(_starts(r.parse_check, "FAIL") for r in rows),
        "Excluded from corpus": sum(_equals(r.lifecycle, "Excluded") for r in rows),
        "Currently in the pilot index (not confirmed eligible)": sum(_equals(r.in_pilot_text, "Yes") for r in rows),
        "Indexed AND eligible": sum(_equals(r.in_pilot_text, "Yes") and _starts(r.lifecycle, "Eligible") for r in rows),
        "Documents with a named open decision (OD-02/03/13)": sum(any(o in r.blocking for o in od) for r in rows),
    }


# ------------------------------------------------------------------------------------------------ the database (read-only)
async def gather_corpus_facts(
    settings: Settings, head: Optional[str] = None, scope: Optional[Iterable[Any]] = None
) -> CorpusFacts:
    """Read-only look at the stored documents. The transaction is declared READ ONLY, so a write would be refused.

    `scope` optionally limits the look to those document IDs (used by tests so they never depend on whatever else a
    shared database holds). The command line never passes it.
    """
    from sqlalchemy import func, select, text
    from sqlalchemy.ext.asyncio import create_async_engine
    from sqlalchemy.pool import NullPool

    from ask_bpl_orchestrator.db.models import Chunk, Document
    from ask_bpl_orchestrator.services.eligibility import retrievable_clause

    facts = CorpusFacts(target=describe_database_target(settings.database_url), allow_test_index=settings.allow_test_index)
    engine = create_async_engine(settings.database_url, poolclass=NullPool, connect_args={"timeout": 5})
    try:
        async with engine.connect() as conn:
            await conn.execute(text("SET TRANSACTION READ ONLY"))
            await conn.execute(text("SELECT 1"))
            try:
                row = (await conn.execute(text("SELECT version_num FROM alembic_version"))).first()
                revision = row[0] if row else None
            except Exception:  # noqa: BLE001 - a missing table means "not migrated"
                await conn.rollback()
                await conn.execute(text("SET TRANSACTION READ ONLY"))
                revision = None
            if revision is None:
                facts.state = "not_migrated"
                await conn.rollback()
                return facts
            if head is not None and revision != head:
                facts.state = "out_of_date"
                await conn.rollback()
                return facts

            ids = None if scope is None else list(scope)
            chunk_counts = select(Chunk.document_id, func.count().label("n")).group_by(Chunk.document_id).subquery()
            query = (
                select(
                    Document.document_id, Document.title, Document.source_path, Document.source_id,
                    Document.eligibility, Document.test_index_member, Document.ingest_status,
                    Document.extraction_status, Document.superseded_by_document_id, chunk_counts.c.n,
                )
                .outerjoin(chunk_counts, chunk_counts.c.document_id == Document.document_id)
                .order_by(Document.source_path, Document.title)
            )
            answerable_query = select(Document.document_id).where(retrievable_clause(settings))
            if ids is not None:
                query = query.where(Document.document_id.in_(ids))
                answerable_query = answerable_query.where(Document.document_id.in_(ids))
            answerable = {str(i) for i in (await conn.execute(answerable_query)).scalars().all()}
            for r in (await conn.execute(query)).all():
                facts.docs.append(
                    DbDoc(
                        document_id=str(r[0]), title=r[1], source_path=r[2], source_id=r[3],
                        eligibility=getattr(r[4], "value", str(r[4])), test_index_member=bool(r[5]),
                        ingest_status=getattr(r[6], "value", str(r[6])), extraction_status=str(r[7]),
                        superseded=r[8] is not None, chunks=int(r[9] or 0),
                        # Retrievable by rule AND holding text. A document with no chunks can answer nothing, so it is
                        # never counted as answerable (the report must not overclaim).
                        answerable=str(r[0]) in answerable and int(r[9] or 0) > 0,
                    )
                )
            facts.state = "checked"
            await conn.rollback()
    except Exception as exc:  # noqa: BLE001 - classify, never print the error's own text
        facts.state = "unreachable"
        facts.error_kind = classify_error(exc)
        facts.docs = []
    finally:
        await engine.dispose()
    return facts


# ------------------------------------------------------------------------------------------------ files, master list
def scan_files(root: Path) -> FilesScan:
    """Hash every file under `root` (recursive). Hidden files and macOS archive litter are skipped, never counted."""
    scan = FilesScan(root=str(root))
    for p in sorted(root.rglob("*")):
        rel_parts = p.relative_to(root).parts
        if not p.is_file() or any(part.startswith(".") or part == "__MACOSX" for part in rel_parts):
            continue
        if p.suffix.lower() not in INGESTIBLE_SUFFIXES:
            scan.ignored += 1
            continue
        digest = hashlib.sha256()
        try:
            with p.open("rb") as fh:
                for block in iter(lambda: fh.read(1 << 20), b""):
                    digest.update(block)
        except OSError:
            scan.unreadable.append(str(Path(*rel_parts)))
            continue
        scan.entries.append(FileEntry(rel="/".join(rel_parts), sha256=digest.hexdigest()))
    return scan


def policy_title_key(text: str, *, is_file: bool = False) -> str:
    """'02 - Materials Selection Policy (04.05) (1).pdf' and 'Materials Selection Policy' both give
    'materials selection policy'. Exact matching only: a near miss stays unmatched (the safe direction)."""
    base = Path(text).name if is_file else text
    if is_file:
        base = re.sub(r"\.[A-Za-z0-9]{2,5}$", "", base)
    base = re.sub(r"\([^)]*\)", " ", base)
    base = re.sub(r"^\s*\d+[A-Za-z]?\s*-\s*", "", base)
    return re.sub(r"[^a-z0-9]+", " ", base.lower()).strip()


def read_master_list(path: Path) -> List[MasterEntry]:
    """A .csv with a `title` column (optional `source_id`, `area`) or a .txt with one title per line."""
    suffix = path.suffix.lower()
    entries: List[MasterEntry] = []
    if suffix == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as fh:
            reader = csv.DictReader(fh)
            names = {(n or "").strip().lower(): n for n in (reader.fieldnames or [])}
            if "title" not in names:
                raise ValueError(f"{path.name}: the first row must have a column named 'title'")
            for r in reader:
                title = _clean(r.get(names["title"]))
                if not title:
                    continue
                sid = _clean(r.get(names["source_id"])) if "source_id" in names else ""
                area = _clean(r.get(names["area"])) if "area" in names else ""
                entries.append(MasterEntry(title, sid or None, area or None))
    elif suffix == ".txt":
        for line in path.read_text(encoding="utf-8-sig").splitlines():
            if line.strip() and not line.lstrip().startswith("#"):
                entries.append(MasterEntry(line.strip()))
    else:
        raise ValueError(f"{path.name}: use a .csv file with a 'title' column, or a .txt file with one title per line")
    if not entries:
        raise ValueError(f"{path.name}: no titles found")
    return entries


# ------------------------------------------------------------------------------------------------ analysis
def expected_state(row: InventoryRow, doc: Optional[DbDoc]) -> tuple[Eligibility, bool]:
    """What the workbook says this document's eligibility and test-index flag should be in the database.

    Mirrors `apply_inventory_to_db`: a row whose parse check claims "Validated" stays blocked unless the stored
    document carries a verified extraction-review record. A test keeps the two in step.
    """
    if row.extraction_validated_claim and (doc is None or doc.extraction_status != ExtractionStatus.validated):
        return Eligibility.blocked, False
    return row.eligibility, row.test_index_member


def _short(text: str, n: int = 46) -> str:
    return text if len(text) <= n else text[: n - 1] + "…"


def _stage(expected: Eligibility, member: bool, doc: Optional[DbDoc], checked: bool) -> str:
    if checked and doc is not None and doc.answerable:
        return STAGE_ANSWERABLE_CONFIRMED if doc.eligibility == Eligibility.eligible.value else STAGE_ANSWERABLE_TEST
    if expected in (Eligibility.excluded, Eligibility.superseded, Eligibility.blocked):
        return expected.value
    if not checked:
        if expected == Eligibility.eligible:
            return "eligible per workbook (database not checked)"
        return "test index per workbook (database not checked)" if member else "authority pending"
    if doc is None:
        return "not ingested"
    if doc.ingest_status != IngestStatus.indexed.value:
        return f"stored, {doc.ingest_status}"
    if doc.chunks == 0:
        return "indexed, no chunks"
    return "indexed, not answerable"


def _row_findings(sr: SourceRow, exp: Eligibility, member: bool, doc: Optional[DbDoc], allow_test_index: bool) -> List[Finding]:
    sid, out = sr.inv.source_id, []

    def add(code: str, text: str, todo: str, severity: str = DISAGREEMENT) -> None:
        out.append(Finding(code, severity, sid, text, todo))

    for w in sr.inv.warnings:
        add("row_warning", w, "Fix the row in the workbook, then run this report again.", NOTE)
    if doc is None:
        if member or exp == Eligibility.eligible:
            add(
                "not_ingested",
                "The workbook expects this document to be able to answer, but nothing is stored for it.",
                "Run askbpl-ingest-batch (the file must be listed in the manifest), then askbpl-apply-inventory.",
            )
        return out
    if doc.answerable and exp in (Eligibility.blocked, Eligibility.excluded, Eligibility.superseded):
        add(
            "answerable_but_not_allowed",
            f"The database lets this document answer, but the workbook says it is {exp.value}.",
            "Run askbpl-apply-inventory --purge-vectors now, then run this report again.",
        )
    elif doc.eligibility != exp.value:
        add(
            "eligibility_differs",
            f"The database says '{doc.eligibility}'; the workbook says '{exp.value}'.",
            "Run askbpl-apply-inventory to bring the database in line with the workbook.",
        )
    if doc.test_index_member != member:
        add(
            "test_index_differs",
            f"The database test-index flag is {doc.test_index_member}; the workbook says {member}.",
            "Run askbpl-apply-inventory.",
        )
    if doc.source_id and doc.source_id != sid:
        add(
            "source_id_differs",
            f"The stored document carries source id {doc.source_id}, but its file name matches {sid}.",
            "Check which file was ingested under this path; run askbpl-apply-inventory once it is settled.",
        )
    if doc.superseded and exp != Eligibility.superseded:
        add(
            "superseded_in_database_only",
            "The database records this document as replaced by a newer one; the workbook does not say so.",
            "Mark the row Superseded in the workbook (see docs/CURATOR_WORKFLOW.md), then run askbpl-apply-inventory.",
        )
    if doc.extraction_status == ExtractionStatus.review_required:
        add(
            "extraction_review_required",
            "Scanned text is stored but no person has validated it, so it can never answer yet.",
            "Follow docs/EXTRACTION_REVIEW.md; do not activate it before the decision is logged.",
            NOTE,
        )
    expects_answer = member or exp == Eligibility.eligible
    if expects_answer and doc.ingest_status != IngestStatus.indexed.value:
        add(
            "not_indexed",
            f"Expected to answer, but its ingest status is '{doc.ingest_status}'.",
            "Run askbpl-ingest-batch again and read its output for a line starting 'failed:'.",
        )
    elif expects_answer and doc.chunks == 0:
        add("no_chunks", "Marked indexed but holds no text chunks.", "Re-ingest it with askbpl-ingest-batch.")
    elif expects_answer and not doc.answerable and not out:
        why = " ALLOW_TEST_INDEX is off, so test-index documents cannot answer." if (member and not allow_test_index) else ""
        add(
            "not_answerable_though_expected",
            "Stored and expected to answer, but the retrieval rule does not allow it." + why,
            "Run askbpl-doctor; check ALLOW_TEST_INDEX and the document's extraction status.",
        )
    return out


def _group_by_key(docs: Sequence[DbDoc]) -> Dict[str, List[DbDoc]]:
    by: Dict[str, List[DbDoc]] = {}
    for d in docs:
        by.setdefault(d.key, []).append(d)
    return by


def _coverage(entries: Sequence[MasterEntry], lines: Sequence[SourceLine], checked: bool) -> tuple[Coverage, List[Finding]]:
    by_id = {ln.source_id: ln for ln in lines}
    by_title: Dict[str, List[SourceLine]] = {}
    for ln in lines:
        by_title.setdefault(policy_title_key(ln.file_name, is_file=True), []).append(ln)
    findings: List[Finding] = []
    seen: Dict[str, MasterEntry] = {}
    unique: List[MasterEntry] = []
    for e in entries:
        k = policy_title_key(e.title)
        if k in seen:
            findings.append(Finding("master_duplicate", NOTE, e.title, "This title appears more than once on the master list; it is counted once.", "Remove the repeat from the list."))
            continue
        seen[k] = e
        unique.append(e)

    items: List[CoverageItem] = []
    matched_ids: set[str] = set()
    for e in unique:
        if e.source_id:
            hit = [by_id[e.source_id]] if e.source_id in by_id else []
            miss = f"source id {e.source_id} is not in the inventory"
        else:
            hit = by_title.get(policy_title_key(e.title), [])
            miss = "no received file has this exact title (add a source_id column to link it)"
        if not hit:
            items.append(CoverageItem(e, "no_file", miss))
        elif len(hit) > 1:
            items.append(CoverageItem(e, "ambiguous", "more than one received file has this title; give its source_id"))
        else:
            ln = hit[0]
            matched_ids.add(ln.source_id)
            if not checked:
                items.append(CoverageItem(e, "received_unchecked", f"{ln.source_id}: {ln.stage}"))
            elif ln.stage.startswith("answerable"):
                items.append(CoverageItem(e, "answerable", f"{ln.source_id}: {ln.stage}"))
            else:
                items.append(CoverageItem(e, "received_not_answerable", f"{ln.source_id}: {ln.stage}"))
    not_on_list = [
        f"{ln.source_id} {_short(ln.file_name, 60)}"
        for ln in lines
        if ln.source_id not in matched_ids and ln.expected not in (Eligibility.excluded.value,)
    ]
    return Coverage(items, not_on_list), findings


def build_report(
    rows: Sequence[SourceRow],
    facts: CorpusFacts,
    *,
    inventory_name: str = "inventory",
    counts_sheet: Optional[Dict[str, Any]] = None,
    files: Optional[FilesScan] = None,
    master: Optional[Sequence[MasterEntry]] = None,
) -> CorpusReport:
    checked = facts.checked
    by_key = _group_by_key(facts.docs) if checked else {}
    findings: List[Finding] = []
    lines: List[SourceLine] = []

    for sr in rows:
        group = by_key.get(sr.inv.key, [])
        doc = group[0] if group else None
        exp, member = expected_state(sr.inv, doc)
        stage = _stage(exp, member, doc, checked)
        row_findings = _row_findings(sr, exp, member, doc, facts.allow_test_index) if checked else [
            Finding("row_warning", NOTE, sr.inv.source_id, w, "Fix the row in the workbook, then run this report again.")
            for w in sr.inv.warnings
        ]
        if len(group) > 1:
            row_findings.append(
                Finding(
                    "duplicate_documents", DISAGREEMENT, sr.inv.source_id,
                    f"{len(group)} stored documents match this one inventory row: " + " | ".join(sorted(d.source_path or d.title for d in group)),
                    "Remove the duplicates (and their vectors) before running askbpl-apply-inventory; it refuses to run while they exist.",
                )
            )
        findings += row_findings
        lines.append(SourceLine(sr.inv.source_id, sr.inv.file_name, sr.source_type, exp.value, member, stage, doc, [f.code for f in row_findings]))

    inventory_keys = {sr.inv.key for sr in rows}
    if checked:
        for key, group in by_key.items():
            if key not in inventory_keys:
                for d in group:
                    findings.append(
                        Finding(
                            "unlisted_document", DISAGREEMENT, d.source_path or d.title,
                            "Stored in the database but not in the inventory workbook"
                            + (" and able to answer." if d.answerable else "; it cannot answer."),
                            "Add an inventory row for it (every received file must be inventoried), then run askbpl-apply-inventory.",
                        )
                    )

    # --- the workbook's own Counts sheet against the rows beneath it
    if counts_sheet is not None:
        for label, recomputed in recompute_workbook_counts(rows).items():
            if label not in counts_sheet:
                continue
            stored = counts_sheet[label]
            if not isinstance(stored, (int, float)):
                findings.append(Finding("counts_sheet_empty", NOTE, label, "The Counts sheet has no stored value for this line (formulas never calculated).", "Open the workbook in Excel and save it."))
            elif int(stored) != recomputed:
                findings.append(
                    Finding(
                        "counts_sheet_differs", DISAGREEMENT, label,
                        f"The workbook's Counts sheet says {int(stored)}; recounting its own rows gives {recomputed}.",
                        "Open the workbook, recalculate (F9), save, and run this report again.",
                    )
                )

    # --- the received files against the inventory (AC-1)
    if files is not None:
        by_name = {sr.inv.key: sr for sr in rows}
        file_keys = set()
        for e in files.entries:
            k = normalize_name(e.rel)
            file_keys.add(k)
            if k not in by_name:
                findings.append(Finding("file_without_row", DISAGREEMENT, e.rel, "This file is in the folder but has no inventory row.", "Add it to the workbook; every received file must be inventoried and dispositioned (AC-1)."))
        for grp in files.duplicate_groups():
            findings.append(Finding("duplicate_content", DISAGREEMENT, grp[0], "These files have identical content: " + " | ".join(grp), "Keep one, record the other as a duplicate in the workbook."))
        for name in files.unreadable:
            findings.append(Finding("file_unreadable", NOTE, name, "This file could not be read, so it was not hashed.", "Check the file's permissions."))
        for sr in rows:
            if sr.inv.key not in file_keys:
                findings.append(Finding("row_without_file", NOTE, sr.inv.source_id, f"No file named like '{_short(sr.inv.file_name, 50)}' is in the folder that was scanned.", "Only a problem if this folder was meant to hold every received file."))

    # --- counts: mutually exclusive buckets that add up to "received"
    elig = [sr.inv.eligibility for sr in rows]
    n = len(rows)
    n_eligible = sum(e == Eligibility.eligible for e in elig)
    n_blocked = sum(e == Eligibility.blocked for e in elig)
    n_excluded = sum(e == Eligibility.excluded for e in elig)
    n_superseded = sum(e == Eligibility.superseded for e in elig)
    n_pending = n - n_eligible - n_blocked - n_excluded - n_superseded
    n_eval_material = sum(_equals(sr.source_type, "Evaluation material") for sr in rows)
    n_member = sum(sr.inv.test_index_member for sr in rows)

    counts = [CountLine("Received (rows in the inventory)", str(n), "AC-1: every received file needs a row")]
    if files is not None:
        counts.append(CountLine("Unique by content", f"{files.unique} of {len(files.entries)} files hashed", "recomputed from the files, not taken from the workbook"))
    else:
        counts.append(CountLine("Unique by content", "not checked here", "the workbook records a manual check; give --files <folder> to recompute"))
    counts += [
        CountLine("  BPL documents (not evaluation material)", str(n - n_eval_material)),
        CountLine("Eligible: BPL-confirmed current", str(n_eligible), "stays 0 until BPL confirms authority (OD-01)"),
        CountLine("Authority pending", str(n_pending), "not blocked, excluded or superseded (the workbook's own line also counts the blocked file)"),
        CountLine("Blocked (parse check failed or extraction blocked)", str(n_blocked), "ING-3"),
        CountLine("Excluded from the corpus", str(n_excluded)),
        CountLine("Superseded", str(n_superseded)),
        CountLine("   those five lines add up to received", str(n_eligible + n_pending + n_blocked + n_excluded + n_superseded)),
        CountLine("Marked as test-index members by the workbook", str(n_member), "answerable only while ALLOW_TEST_INDEX is on"),
        CountLine("Evaluated", "not tracked", NOT_EVALUATED_NOTE),
    ]
    answerable_now: Optional[int] = None
    if checked:
        docs = facts.docs
        indexed = [d for d in docs if d.ingest_status == IngestStatus.indexed.value and d.chunks > 0]
        answerable = [d for d in docs if d.answerable]
        confirmed = [d for d in answerable if d.eligibility == Eligibility.eligible.value]
        answerable_now = len(answerable)
        counts += [
            CountLine("Stored in the database", str(len(docs)), facts.target),
            CountLine("Indexed (stored, indexed, has text chunks)", str(len(indexed))),
            CountLine("Answerable now", str(len(answerable)), "ALLOW_TEST_INDEX is " + ("on" if facts.allow_test_index else "off")),
            CountLine("   of which BPL-confirmed (eligible)", str(len(confirmed))),
            CountLine("   of which test index only", str(len(answerable) - len(confirmed))),
            CountLine("Indexed AND eligible", str(sum(d.eligibility == Eligibility.eligible.value for d in indexed)), "cannot be claimed until BPL confirms"),
        ]
    else:
        counts.append(CountLine("Stored, indexed, answerable now", "not checked", _db_state_words(facts)))

    coverage: Optional[Coverage] = None
    if master is not None:
        coverage, cov_findings = _coverage(master, lines, checked)
        findings += cov_findings

    headline = _headline(facts, answerable_now, n_member, n_eligible)
    return CorpusReport(inventory_name, facts, counts, lines, findings, coverage, headline, n, answerable_now, n_eligible)


def _db_state_words(facts: CorpusFacts) -> str:
    return {
        "skipped": "the database was not read (--no-db)",
        "unreachable": f"the database could not be reached ({facts.error_kind or 'unknown'}); run askbpl-doctor",
        "not_migrated": "the database has no tables yet; run alembic upgrade head, then askbpl-seed",
        "out_of_date": "the database is older than this code; run alembic upgrade head",
    }.get(facts.state, "")


def _headline(facts: CorpusFacts, answerable: Optional[int], members: int, eligible: int) -> str:
    if not facts.checked:
        return (
            f"The database was not checked: {_db_state_words(facts)}. This report therefore cannot say what can answer now. "
            f"The workbook marks {members} as test-index members. BPL-confirmed (eligible) documents: {eligible}."
        )
    if not answerable:
        return f"No document can answer questions now. BPL-confirmed (eligible) documents: {eligible}."
    test_only = answerable - sum(1 for d in facts.docs if d.answerable and d.eligibility == Eligibility.eligible.value)
    tail = (
        f"All {answerable} are test-index documents that BPL has not confirmed as current."
        if test_only == answerable
        else f"{test_only} of them are test-index documents that BPL has not confirmed as current."
    )
    plural = "document" if answerable == 1 else "documents"
    return f"{answerable} {plural} can answer questions now. {tail} BPL-confirmed (eligible) documents: {eligible}."


# ------------------------------------------------------------------------------------------------ rendering
def render(report: CorpusReport, *, summary_only: bool = False) -> str:
    out: List[str] = []
    w = out.append
    w("Ask BPL corpus report (read-only)")
    w(f"Inventory: {report.inventory_name}")
    w("")
    w(report.headline)
    w("")
    w("COUNTS (PRD ING-13)")
    width = max(len(c.label) for c in report.counts)
    for c in report.counts:
        note = f"   {c.note}" if c.note else ""
        w(f"  {c.label.ljust(width)}  {c.value}{note}")
    w("")

    w("COVERAGE OF THE POLICY SET")
    if report.coverage is None:
        w("  Uncovered policy areas: unknown. No master list of policies was supplied, so this report cannot say what is")
        w("  missing. The counts above are documents received, not coverage of the full policy set. Give --master-list")
        w("  <file> (a .csv with a 'title' column, optional 'source_id' and 'area') once BPL's list exists.")
    else:
        _render_coverage(report, w)
    w("")

    dis = report.disagreements
    notes = [f for f in report.findings if f.severity == NOTE]
    if dis:
        w(f"THINGS TO RECONCILE ({len(dis)})")
    elif report.facts.checked:
        w("THINGS TO RECONCILE (none found)")
    else:
        w("THINGS TO RECONCILE (none found in the workbook alone)")
    if not report.facts.checked:
        w(f"  The workbook-versus-database checks were not run: {_db_state_words(report.facts)}.")
    for f in dis:
        w(f"  [{f.code}] {f.subject}: {f.text}")
        if f.todo:
            w(f"      do: {f.todo}")
    if notes:
        w("")
        w(f"NOTES ({len(notes)})")
        for f in notes:
            w(f"  [{f.code}] {f.subject}: {f.text}")
    w("")

    if not summary_only:
        w("PER SOURCE")
        w(f"  {'source':8} {'file':48} {'workbook says':23} {'stage':36} flags")
        for ln in report.lines:
            flags = ",".join(ln.finding_codes)
            w(f"  {ln.source_id:8} {_short(ln.file_name, 48):48} {ln.expected + ('+test' if ln.expected_test_index else ''):23} {ln.stage:36} {flags}")
        stored = [ln for ln in report.lines if ln.doc is not None]
        if stored:
            w("")
            w("STORED DOCUMENTS (use the full id with askbpl-ingest --supersedes)")
            for ln in stored:
                w(f"  {ln.source_id:8} {ln.doc.document_id}  chunks={ln.doc.chunks}  {ln.doc.ingest_status}")
        w("")
    w("WHAT THIS REPORT CANNOT TELL YOU")
    for s in CANNOT_TELL:
        w(f"  - {s}")
    return "\n".join(out) + "\n"


def _render_coverage(report: CorpusReport, w: Any) -> None:
    cov = report.coverage
    assert cov is not None
    total = len(cov.items)
    checked = report.facts.checked
    w(f"  Policies on the master list: {total}")
    if checked:
        w(f"  Have an answerable document now: {cov.count('answerable')} of {total} (test index; BPL-confirmed documents: {report.eligible})")
        w(f"  Received but cannot answer yet:  {cov.count('received_not_answerable')}")
    else:
        w(f"  Matched to a received file:      {cov.count('received_unchecked')} (database not checked, so answerable is unknown)")
    w(f"  No file matched or ambiguous:    {cov.count('no_file') + cov.count('ambiguous')}")
    gaps = [i for i in cov.items if i.status != "answerable"]
    if gaps:
        by_area: Dict[str, List[CoverageItem]] = {}
        for i in gaps:
            by_area.setdefault(i.entry.area or "(no area given)", []).append(i)
        w("  No answerable document for:")
        for area in sorted(by_area):
            w(f"    {area}")
            for i in by_area[area]:
                w(f"      - {i.entry.title}: {i.detail}")
    if cov.not_on_list:
        shown = cov.not_on_list[:15]
        w(f"  Received but not matched to any master-list title ({len(cov.not_on_list)}):")
        for s in shown:
            w(f"      - {s}")
        if len(cov.not_on_list) > len(shown):
            w(f"      ... and {len(cov.not_on_list) - len(shown)} more")


def _csv_safe(value: Any) -> str:
    text = "" if value is None else str(value)
    return "'" + text if text[:1] in ("=", "+", "-", "@") else text


def write_csv(report: CorpusReport, path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8") as fh:
        wr = csv.writer(fh)
        wr.writerow(["source_id", "file_name", "source_type", "workbook_eligibility", "workbook_test_index", "stage", "document_id", "ingest_status", "chunks", "answerable_now", "flags"])
        for ln in report.lines:
            d = ln.doc
            wr.writerow([
                _csv_safe(ln.source_id), _csv_safe(ln.file_name), _csv_safe(ln.source_type), ln.expected, ln.expected_test_index, ln.stage,
                d.document_id if d else "", d.ingest_status if d else "", d.chunks if d else "", d.answerable if d else "", ";".join(ln.finding_codes),
            ])


# ------------------------------------------------------------------------------------------------ command line
def _find_inventory(arg: str) -> Optional[Path]:
    p = Path(arg)
    options = [p] if p.is_absolute() else [Path.cwd() / p, Path(__file__).resolve().parents[1] / p]
    return next((o for o in options if o.is_file()), None)


def main(argv: Optional[Sequence[str]] = None) -> int:
    from ask_bpl_orchestrator.config import get_settings

    settings = get_settings()
    ap = argparse.ArgumentParser(description="Counts of the policy corpus: received, eligible, blocked, excluded, indexed, answerable (read-only)")
    ap.add_argument("--inventory", default=settings.source_inventory_path, help="the source inventory workbook (default: SOURCE_INVENTORY_PATH)")
    ap.add_argument("--no-db", action="store_true", help="read only the workbook; do not look at the database")
    ap.add_argument("--files", metavar="FOLDER", help="hash the received policy files in this folder (unique by content; files never inventoried)")
    ap.add_argument("--master-list", metavar="FILE", help="BPL's list of policies (.csv with a 'title' column, or .txt): says which have no answerable document")
    ap.add_argument("--csv", metavar="FILE", help="also write the per-source table here (the only file this command writes)")
    ap.add_argument("--summary-only", action="store_true", help="skip the per-source table")
    ap.add_argument("--strict", action="store_true", help="exit 1 if anything needs reconciling, or the database was asked for and could not be read")
    args = ap.parse_args(argv)

    inv_path = _find_inventory(args.inventory)
    if inv_path is None:
        print(f"Cannot find the inventory workbook '{args.inventory}'. Give --inventory <file>, or set SOURCE_INVENTORY_PATH.", file=sys.stderr)  # noqa: T201
        return 2
    try:
        rows = read_sources(inv_path)
        counts_sheet = read_counts_sheet(inv_path)
        files = None
        if args.files:
            root = Path(args.files)
            if not root.is_dir():
                raise ValueError(f"--files: '{args.files}' is not a folder")
            files = scan_files(root)
        master = read_master_list(Path(args.master_list)) if args.master_list else None
    except (ValueError, OSError, KeyError) as exc:
        print(f"Cannot read the input: {exc}", file=sys.stderr)  # noqa: T201
        return 2

    facts = CorpusFacts() if args.no_db else asyncio.run(gather_corpus_facts(settings, alembic_head()))
    report = build_report(rows, facts, inventory_name=inv_path.name, counts_sheet=counts_sheet, files=files, master=master)
    print(render(report, summary_only=args.summary_only), end="")  # noqa: T201
    if args.csv:
        write_csv(report, Path(args.csv))
        print(f"Wrote {args.csv}")  # noqa: T201
    if args.strict:
        unreadable = (not args.no_db) and not facts.checked
        if report.disagreements or unreadable:
            return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
